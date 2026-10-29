"""Read inputs from workbook sheets (DataFrames) and write results to Excel.

Input sheets (same layout in the .xlsm front end and in a plain .xlsx used by the CLI):
  Config       key | value                            (reporting_ccy: currency of all exposure outputs)
  Market       underlying | spot | div_yield | is_index | flat_vol | bbg_ticker | ccy
  Curve        tenor | zero_rate | ccy                (continuously compounded, decimal; one curve per ccy)
  FX           pair | spot                            (optional; e.g. USDCAD | 1.40 = CAD per 1 USD)
  VolSurface   underlying | tenor | moneyness | vol    (optional; long format)
  Correlation  underlying1 | underlying2 | rho         (optional)
  NettingSets  netting_set | counterparty | collateralized | threshold_cp | threshold_bank | mta |
               ia_received | ia_posted | ia_posted_segregated | mpor_days | rounding | csa_ccy
Blank ccy / csa_ccy / Curve ccy mean the reporting currency.
  Trades       trade_id | netting_set | leg_id | leg_type | underlying | direction | shares |
               start_date | maturity | equity_reset_freq | funding_reset_freq | funding_spread |
               notional_basis | avg_price | initial_mtm | div_passthrough |
               option_type | strike | option_maturity | vol_override | ovme_price
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from . import dates as dt
from .market import MarketData, Underlying, VolSurface, ZeroCurve
from .model import RunConfig, build_netting_sets

INPUT_SHEETS = ["Config", "Market", "Curve", "FX", "VolSurface", "Correlation", "NettingSets", "Trades"]


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df.loc[:, [c for c in df.columns if c and not c.startswith("Unnamed") and c != "None"]]
    return df.dropna(how="all")


def _num(x, what: str) -> float:
    try:
        v = float(x)
    except (TypeError, ValueError):
        raise ValueError(f"{what}: '{x}' is not a number (Bloomberg still loading or #N/A?)") from None
    if np.isnan(v):
        raise ValueError(f"{what}: missing value")
    return v


def _truthy(x) -> bool:
    return str(x).strip().upper() in ("Y", "YES", "TRUE", "1", "1.0")


def _records(df: pd.DataFrame) -> list[dict]:
    recs = []
    for r in df.to_dict("records"):
        recs.append({k: (None if (isinstance(v, float) and np.isnan(v)) else v) for k, v in r.items()})
    return recs


def _ccy(x) -> str:
    return "" if x is None or (isinstance(x, float) and np.isnan(x)) else str(x).strip().upper()


def parse_fx(df: pd.DataFrame | None, rep: str) -> dict[str, float]:
    """FX rows 'BASEQUOTE | rate' (1 BASE = rate QUOTE; a Bloomberg ticker such as 'USDCAD Curncy'
    works too) -> reporting-currency units per 1 unit of each currency. Crosses are chained."""
    pairs = {}
    if df is not None:
        for r in _records(_clean(df)):
            p = "".join(ch for ch in str(r.get("pair") or "").upper() if ch.isalpha())[:6]
            if len(p) == 6 and r.get("spot") is not None:
                pairs[(p[:3], p[3:])] = _num(r["spot"], f"FX {p}")
    rate = {rep: 1.0}
    while pairs:
        done = [k for k in pairs if k[0] in rate or k[1] in rate]
        if not done:
            break
        for b, q in done:
            r = pairs.pop((b, q))
            if q in rate and b not in rate:
                rate[b] = r * rate[q]
            elif b in rate and q not in rate:
                rate[q] = rate[b] / r
    rate.pop(rep)
    return rate


def parse_config(df: pd.DataFrame) -> dict:
    df = _clean(df)
    kv = {str(k).strip().lower(): v for k, v in zip(df.iloc[:, 0], df.iloc[:, 1]) if str(k).strip()}
    return kv


def parse_inputs(sheets: dict[str, pd.DataFrame]):
    """Return (market, netting_sets, run_config, warnings)."""
    warnings: list[str] = []
    kv = parse_config(sheets["Config"])
    val = dt.to_date(kv["valuation_date"])
    market_df = _clean(sheets["Market"])
    curve_df = _clean(sheets["Curve"])
    ns_rows = _records(_clean(sheets["NettingSets"])) if "NettingSets" in sheets else []

    # reporting currency: Config value, else the single currency used everywhere (blank = single-ccy run)
    rep = _ccy(kv.get("reporting_ccy"))
    if not rep:
        used = {_ccy(c) for df_, col in ((market_df, "ccy"), (curve_df, "ccy")) if col in df_
                for c in df_[col]} | {_ccy(r.get("csa_ccy")) for r in ns_rows}
        used.discard("")
        if len(used) > 1:
            raise ValueError(f"Several currencies used ({', '.join(sorted(used))}): set reporting_ccy in Config")
        rep = used.pop() if used else ""

    curve_df = curve_df[curve_df["tenor"].notna() & (curve_df["tenor"].astype(str).str.strip() != "")]
    cc = curve_df["ccy"].map(_ccy).replace("", rep) if "ccy" in curve_df else pd.Series(rep, curve_df.index)
    curves: dict[str, ZeroCurve] = {}
    for c in dict.fromkeys(cc):
        sub = curve_df[cc == c]
        ts = np.array([dt.yf(val, dt.date_or_tenor(str(t).strip().upper(), val)) for t in sub["tenor"]])
        rs = np.array([_num(r, f"Curve {c} {t}") for t, r in zip(sub["tenor"], sub["zero_rate"])])
        order = np.argsort(ts)
        curves[c] = ZeroCurve(ts[order], rs[order])
    if rep not in curves:
        raise ValueError(f"No discount curve for the reporting currency {rep}: add Curve rows with ccy = {rep}")
    curve = curves.pop(rep)

    vs = _clean(sheets["VolSurface"]) if "VolSurface" in sheets and sheets["VolSurface"] is not None else None
    unds: dict[str, Underlying] = {}
    for r in _records(market_df):
        name = str(r["underlying"]).strip()
        if not name or name == "None":
            continue
        surf = None
        if vs is not None and len(vs):
            sub = vs[vs["underlying"].astype(str).str.strip() == name]
            sub = sub[pd.to_numeric(sub["vol"], errors="coerce").notna()]   # blank = no BBG data, no override
            if len(sub):
                sub = sub.assign(T=[dt.tenor_years(str(t).strip().upper()) if dt.is_tenor(str(t).strip())
                                    else float(t) for t in sub["tenor"]],
                                 m=sub["moneyness"].astype(float), v=[_num(v, f"Vol {name}") for v in sub["vol"]])
                piv = sub.pivot_table(index="T", columns="m", values="v")
                if piv.isna().any().any():
                    piv = piv.interpolate(axis=1, limit_direction="both")
                surf = VolSurface(piv.index.values.astype(float), piv.columns.values.astype(float), piv.values)
        if surf is None:
            surf = VolSurface.flat(_num(r.get("flat_vol"), f"Market {name} flat_vol"))
        unds[name] = Underlying(name, _num(r["spot"], f"Market {name} spot"),
                                0.0 if r.get("div_yield") is None else _num(r["div_yield"], f"{name} div_yield"),
                                _truthy(r.get("is_index", "N")), surf, str(r.get("bbg_ticker") or ""),
                                _ccy(r.get("ccy")) or rep)
    corr = {}
    if "Correlation" in sheets and sheets["Correlation"] is not None:
        for r in _records(_clean(sheets["Correlation"])):
            if r.get("underlying1") and r.get("underlying2"):
                corr[(str(r["underlying1"]).strip(), str(r["underlying2"]).strip())] = float(r["rho"])
    mkt = MarketData(val, curve, unds, corr, rep, curves, parse_fx(sheets.get("FX"), rep))

    trade_rows = _records(_clean(sheets["Trades"]))
    nss = build_netting_sets(trade_rows, ns_rows, mkt, warnings)
    for ns in nss:
        if not ns.csa.ccy:
            ns.csa.ccy = rep
    for c in {u.ccy for u in unds.values()} | {ns.csa.ccy for ns in nss}:
        mkt.curve_for(c)          # fail now, with a clear message, if a curve or FX rate is missing
        mkt.fx(c, rep)
    horizon = val + dt.timedelta(days=int(max(ns.maturity for ns in nss) * 365) + 1)
    cfg = RunConfig(
        valuation_date=val,
        buckets=dt.parse_buckets(kv.get("buckets", "1M,3M,6M,1Y,2Y,3Y,5Y"), val, horizon),
        n_paths=int(float(kv.get("n_paths", 10000))),
        seed=int(float(kv.get("seed", 42))),
        pfe_quantile=float(kv.get("pfe_quantile", 0.95)),
        fine_step=str(kv.get("fine_step", "1W")).strip().upper(),
        output_mode=str(kv.get("output_mode", "AVERAGE")).strip().upper(),
        antithetic=_truthy(kv.get("antithetic", "Y")),
        epe_horizon=str(kv.get("epe_horizon", "1Y")).strip().upper(),
        analytic_nodes=int(float(kv.get("analytic_nodes", 201))),
    )
    return mkt, nss, cfg, warnings


def read_workbook(path: str | Path) -> dict[str, pd.DataFrame]:
    xl = pd.ExcelFile(path)
    out = {}
    for s in INPUT_SHEETS:
        if s in xl.sheet_names:
            out[s] = xl.parse(s, header=0 if s != "Config" else 0)
    missing = {"Config", "Market", "Curve", "Trades"} - set(out)
    if missing:
        raise ValueError(f"{path}: missing sheets {sorted(missing)}")
    return out


def write_results(results, path: str | Path, inputs: dict[str, pd.DataFrame] | None = None) -> Path:
    from openpyxl.chart import LineChart, Reference, Series
    from openpyxl.utils import get_column_letter

    path = Path(path)
    tables = results.tables()
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in tables.items():
            df.to_excel(xw, sheet_name=name, index=False)
        if inputs:
            for name, df in inputs.items():
                df.to_excel(xw, sheet_name=f"in_{name}"[:31], index=False)
        wb = xw.book
        for ws in wb.worksheets:
            for col in ws.columns:
                w = max(len(str(c.value)) if c.value is not None else 0 for c in col[:50])
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, w + 2), 40)
            ws.freeze_panes = "A2"
        # charts: one per netting set from EE_NEE
        ws = wb["EE_NEE"]
        cs = wb.create_sheet("Charts", 1)
        df = tables["EE_NEE"]
        cols = list(df.columns)
        anchor_row = 1
        for ns, g in df.groupby("netting_set", sort=False):
            r0, r1 = g.index.min() + 2, g.index.max() + 2
            ch = LineChart()
            ch.title = f"{ns}: EE / NEE / PFE"
            ch.y_axis.title = "Exposure"
            ch.x_axis.title = "Bucket"
            for c in ("EE", "NEE", "PFE", "EE_uncoll", "SACCR_EAD_mean"):
                ci = cols.index(c) + 1
                ch.series.append(Series(Reference(ws, min_col=ci, min_row=r0, max_row=r1), title=c))
            ci = cols.index("bucket_date") + 1
            ch.set_categories(Reference(ws, min_col=ci, min_row=r0, max_row=r1))
            ch.height, ch.width = 8, 18
            cs.add_chart(ch, f"A{anchor_row}")
            anchor_row += 17
    return path
